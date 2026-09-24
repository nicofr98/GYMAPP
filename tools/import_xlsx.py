#!/usr/bin/env python3
"""Convert the Upper/Lower Excel sheet into GymApp data.

Usage:
    python3 tools/import_xlsx.py <workbook.xlsx> [backup.json]

Always writes js/default-program.js: the program for each person, with no
logged numbers, so it is safe to commit to a public repo.

With a second argument it also writes a full backup (program + the weeks
already logged). Import that file on the phone from Settings. Do not commit it.

Requires: pip install openpyxl
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
WEEKS = 8
DELOAD_WEEK = 8
FIRST_WEEK_COL = 9  # column I holds week 1 reps, J week 1 weight, K week 2 reps...

DAY_NAMES = {
    1: "Lower A (Strength)",
    2: "Upper A (Strength)",
    3: "Lower B (Glutes + Hamstrings)",
    4: "Upper B (Hypertrophy)",
}

NOTES_EN = {
    "Serie pesada tras calentar. Profundidad mínima: paralela.":
        "Heavy set after warming up. Minimum depth: parallel.",
    "~10% menos peso que la top set.": "~10% less weight than the top set.",
    "Cadera atrás hasta el máximo estiramiento del femoral. Espalda neutra.":
        "Hips back until max hamstring stretch. Neutral back.",
    "Inclina el torso adelante para estirar más el femoral. Última serie al fallo.":
        "Lean the torso forward to stretch the hamstrings more. Last set to failure.",
    "Pausa de 1 s arriba. Última serie al fallo.": "1 s pause at the top. Last set to failure.",
    "Torso inclinado adelante = más glúteo. Opción: MACHINE HIP THRUST para glúteo mayor directo.":
        "Torso leaning forward = more glutes. Option: MACHINE HIP THRUST for direct glute max work.",
    "Pausa de 1-2 s abajo en el estiramiento.": "1-2 s pause at the bottom in the stretch.",
    "Escápulas retraídas, toca el pecho en cada rep.": "Shoulder blades retracted, touch the chest every rep.",
    "Torso a ~45°, sin impulso con la cadera.": "Torso at ~45°, no hip drive.",
    "Lleva los codos hacia las costillas.": "Drive the elbows toward the ribs.",
    "Baja hasta que las mancuernas queden a la altura de las orejas.":
        "Lower until the dumbbells are at ear height.",
    "Superserie con A2.": "Superset with A2.",
    "Estira bien el tríceps abajo.": "Get a full triceps stretch at the bottom.",
    "Sube hacia afuera, no hacia arriba. Última serie al fallo.":
        "Raise out, not up. Last set to failure.",
    "Pausa de 1 s arriba, barbilla hacia el pecho.": "1 s pause at the top, chin tucked.",
    "Paso largo y torso inclinado para más glúteo.": "Long stride and leaning torso for more glutes.",
    "Última serie al fallo.": "Last set to failure.",
    "Controla la bajada.": "Control the lowering.",
    "Pausa de 1-2 s abajo.": "1-2 s pause at the bottom.",
    "Banco a 30-45°.": "Bench at 30-45°.",
    "Rango completo, estira arriba.": "Full range, stretch at the top.",
    "Aprieta escápulas 1 s.": "Squeeze the shoulder blades for 1 s.",
    "Enfócate en el estiramiento. Última serie al fallo.": "Focus on the stretch. Last set to failure.",
    "Baja la barra detrás de la cabeza para más estiramiento.":
        "Lower the bar behind the head for more stretch.",
}

# Personal remarks after "//" in a note decide how the weight column is read.
WEIGHT_TYPE_REMARKS = {
    "peso propio": "bodyweight",
    "El peso es de asistencia": "assisted",
}

# Rows whose values were typed N columns to the right of where they belong:
# (sheet, row) -> N.
SHIFTED_ROWS = {
    ("Nico", 15): 1,  # Hip abduction: week 1 reps is in J instead of I.
}


def translate_reps(s):
    return s.replace("c/pierna", "/leg")


def translate_rest(s):
    return s.replace(" SEG", " s").replace(" MIN", " min")


def increment_for(name):
    n = name.upper()
    if "DUMBBELL" in n or "HAMMER" in n or "BULGARIAN" in n:
        return 2
    if "BARBELL" in n or "SQUAT" in n or "ROMANIAN" in n or "EZ BAR" in n or "CABLE" in n:
        return 2.5
    return 5


def num(v):
    """Numeric cell value, or None for blanks and placeholders like '.'."""
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    return None


def parse_sheet(ws, person_id):
    block_id = f"{person_id}-b1"
    days = []
    sessions = {}
    day = None
    day_no = 0
    for row in range(1, ws.max_row + 1):
        c = ws.cell(row, 3).value
        if isinstance(c, str) and c.startswith("DÍA"):
            n = int(re.search(r"DÍA (\d+)", c).group(1))
            day_no = n
            day = {"id": f"{block_id}-d{n}", "name": DAY_NAMES.get(n, c), "exercises": []}
            days.append(day)
            continue
        if day is None or not isinstance(c, str) or c == "EXERCISE":
            continue

        name = c.strip()
        superset = ""
        m = re.match(r"^([A-Z]\d)\.\s*(.+)$", name)
        if m:
            superset, name = m.group(1), m.group(2)

        raw_notes = (ws.cell(row, 26).value or "").strip()
        note, _, remark = (p.strip() for p in raw_notes.partition("//"))
        weight_type = WEIGHT_TYPE_REMARKS.get(remark, "load")

        ex_id = f"{day['id']}-e{len(day['exercises']) + 1}"
        sets = int(ws.cell(row, 5).value or 1)
        day["exercises"].append({
            "id": ex_id,
            "name": name,
            "superset": superset,
            "warmup": str(ws.cell(row, 4).value or "0"),
            "sets": sets,
            "reps": translate_reps(str(ws.cell(row, 6).value or "")),
            "rpe": str(ws.cell(row, 7).value or ""),
            "rest": translate_rest(str(ws.cell(row, 8).value or "")),
            "sub": (ws.cell(row, 25).value or "").strip(),
            "notes": NOTES_EN.get(note, note),
            "weightType": weight_type,
            "increment": increment_for(name),
        })

        shift = SHIFTED_ROWS.get((ws.title, row), 0)
        for week in range(1, WEEKS + 1):
            col = FIRST_WEEK_COL + 2 * (week - 1) + shift
            reps = num(ws.cell(row, col).value)
            raw_weight = ws.cell(row, col + 1).value
            weight = num(raw_weight)
            if reps is None and weight is None:
                continue
            entry = {
                "name": name,
                "sets": [{"reps": reps, "weight": weight, "done": True} for _ in range(sets)],
            }
            if reps is None:
                entry["flag"] = "Imported from Excel without reps. Check this."
            elif weight is None and weight_type == "load" and raw_weight != ".":
                entry["flag"] = "Imported from Excel without weight. Check this."
            key = (week, day["id"])
            sessions.setdefault(key, {
                "id": f"imp-{person_id}-w{week}-d{day_no}",
                "personId": person_id,
                "blockId": block_id,
                "week": week,
                "dayId": day["id"],
                "date": None,
                "imported": True,
                "entries": {},
            })["entries"][ex_id] = entry

    block = {
        "id": block_id,
        "personId": person_id,
        "name": "Upper/Lower block 1",
        "weeks": WEEKS,
        "deloadWeek": DELOAD_WEEK,
        "createdAt": date.today().isoformat(),
        "days": days,
    }
    return block, list(sessions.values())


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1])
    people, blocks, sessions = [], [], []
    for ws in wb.worksheets:
        person_id = ws.title.strip().lower()
        block, person_sessions = parse_sheet(ws, person_id)
        people.append({"id": person_id, "name": ws.title.strip(), "activeBlockId": block["id"]})
        blocks.append(block)
        sessions.extend(person_sessions)

    state = {"version": 1, "activePersonId": people[0]["id"], "people": people, "blocks": blocks, "sessions": []}
    out = ROOT / "js" / "default-program.js"
    out.write_text(
        "// Generated by tools/import_xlsx.py. Program only, no logged numbers.\n"
        f"export const DEFAULT_STATE = {json.dumps(state, indent=1, ensure_ascii=False)};\n",
        encoding="utf-8",
    )
    print(f"wrote {out}")

    if len(sys.argv) > 2:
        backup = dict(state, sessions=sessions)
        Path(sys.argv[2]).write_text(json.dumps(backup, ensure_ascii=False), encoding="utf-8")
        flagged = [(s["personId"], s["week"], e["name"], e["flag"])
                   for s in sessions for e in s["entries"].values() if "flag" in e]
        print(f"wrote {sys.argv[2]}: {len(sessions)} sessions")
        for f in flagged:
            print("  flagged:", *f)


if __name__ == "__main__":
    main()
