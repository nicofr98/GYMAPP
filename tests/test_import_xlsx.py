"""Tests for tools/import_xlsx.py, using a small workbook built on the fly."""
import sys
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
import import_xlsx  # noqa: E402

HEADER = ["EXERCISE", "WARM-UP", "SETS", "REPS", "RPE", "REST"]


def make_sheet(title, rows):
    """rows: (row_number, [C..H values], {week: (reps, weight)}, sub, notes)."""
    ws = openpyxl.Workbook().active
    ws.title = title
    ws.cell(8, 3, "DÍA 1 – LOWER A (FUERZA)")
    for i, v in enumerate(HEADER):
        ws.cell(9, 3 + i, v)
    for row, rx, weeks, sub, notes in rows:
        for i, v in enumerate(rx):
            ws.cell(row, 3 + i, v)
        for week, (reps, weight) in weeks.items():
            col = import_xlsx.FIRST_WEEK_COL + 2 * (week - 1)
            ws.cell(row, col, reps)
            ws.cell(row, col + 1, weight)
        ws.cell(row, 25, sub)
        ws.cell(row, 26, notes)
    return ws


class ParseSheetTest(unittest.TestCase):
    def setUp(self):
        self.ws = make_sheet("Test", [
            (10, ["BACK SQUAT (TOP SET)", "3-4", 1, "4-6", "8-9", "3 MIN"], {1: (5, 100), 2: (6, 100)},
             "HACK SQUAT", "Serie pesada tras calentar. Profundidad mínima: paralela."),
            (11, ["BULGARIAN SPLIT SQUAT", "1", 3, "8-10 c/pierna", "8-9", "90 SEG"], {1: (None, 20)},
             "DUMBBELL WALKING LUNGE", "Nota nueva sin traducir"),
            (12, ["LAT PULLDOWN", "1", 3, "8-12", "8-9", "2 MIN"], {1: (8, ".")},
             "PULL-UP", "Lleva los codos hacia las costillas. // peso propio"),
            (13, ["A1. INCLINE DUMBBELL CURL", "0", 2, "10-12", "9-10", "0 SEG → A2"], {1: (10, None)},
             "BAYESIAN CABLE CURL", "Superserie con A2."),
            (14, ["MACHINE SEATED HIP ABDUCTION", "0", 2, "12-15", "10", "60-90 SEG"], {},
             "CABLE HIP ABDUCTION", ""),
        ])
        self.block, self.sessions = import_xlsx.parse_sheet(self.ws, "test")
        self.ex = self.block["days"][0]["exercises"]

    def test_block_shape(self):
        self.assertEqual(self.block["id"], "test-b1")
        self.assertEqual(self.block["weeks"], 8)
        self.assertEqual(self.block["deloadWeek"], 8)
        self.assertEqual(len(self.block["days"]), 1)
        self.assertEqual(self.block["days"][0]["name"], "Lower A (Strength)")
        self.assertEqual(len(self.ex), 5)

    def test_prescription_is_translated(self):
        squat, bulgarian = self.ex[0], self.ex[1]
        self.assertEqual(squat["rest"], "3 min")
        self.assertEqual(squat["notes"], "Heavy set after warming up. Minimum depth: parallel.")
        self.assertEqual(bulgarian["reps"], "8-10 /leg")
        self.assertEqual(bulgarian["rest"], "90 s")
        self.assertEqual(bulgarian["notes"], "Nota nueva sin traducir")  # unknown notes are kept

    def test_superset_prefix_is_split_off(self):
        curl = self.ex[3]
        self.assertEqual(curl["superset"], "A1")
        self.assertEqual(curl["name"], "INCLINE DUMBBELL CURL")

    def test_weight_type_from_personal_remark(self):
        self.assertEqual(self.ex[0]["weightType"], "load")
        self.assertEqual(self.ex[2]["weightType"], "bodyweight")
        self.assertEqual(self.ex[2]["notes"], "Drive the elbows toward the ribs.")

    def test_increments(self):
        self.assertEqual(self.ex[0]["increment"], 2.5)  # barbell squat
        self.assertEqual(self.ex[1]["increment"], 2)    # bulgarian (dumbbells)
        self.assertEqual(self.ex[3]["increment"], 2)    # dumbbell curl
        self.assertEqual(self.ex[4]["increment"], 5)    # machine

    def test_sessions_per_week_with_one_value_per_set(self):
        weeks = sorted(s["week"] for s in self.sessions)
        self.assertEqual(weeks, [1, 2])
        w1 = next(s for s in self.sessions if s["week"] == 1)
        self.assertEqual(w1["id"], "imp-test-w1-d1")
        self.assertTrue(w1["imported"])
        squat = w1["entries"]["test-b1-d1-e1"]
        self.assertEqual(squat["sets"], [{"reps": 5, "weight": 100, "done": True}])
        self.assertNotIn("test-b1-d1-e5", w1["entries"])  # nothing logged

    def test_missing_values_are_flagged_not_guessed(self):
        w1 = next(s for s in self.sessions if s["week"] == 1)
        bulgarian = w1["entries"]["test-b1-d1-e2"]
        self.assertEqual(len(bulgarian["sets"]), 3)
        self.assertIsNone(bulgarian["sets"][0]["reps"])
        self.assertIn("without reps", bulgarian["flag"])
        self.assertIn("without weight", w1["entries"]["test-b1-d1-e4"]["flag"])

    def test_bodyweight_dot_is_not_flagged(self):
        w1 = next(s for s in self.sessions if s["week"] == 1)
        pulldown = w1["entries"]["test-b1-d1-e3"]
        self.assertIsNone(pulldown["sets"][0]["weight"])
        self.assertNotIn("flag", pulldown)


class ShiftedRowTest(unittest.TestCase):
    def test_nico_hip_abduction_is_realigned(self):
        ws = make_sheet("Nico", [
            (15, ["MACHINE SEATED HIP ABDUCTION", "0", 2, "12-15", "10", "60-90 SEG"], {}, "", ""),
        ])
        # As in the real sheet: J15=12, K15=85, L15=15, M15=85 (one column right).
        for col, v in zip("JKLM", [12, 85, 15, 85]):
            ws[f"{col}15"] = v
        _, sessions = import_xlsx.parse_sheet(ws, "nico")
        by_week = {s["week"]: s["entries"]["nico-b1-d1-e1"]["sets"][0] for s in sessions}
        self.assertEqual(by_week[1], {"reps": 12, "weight": 85, "done": True})
        self.assertEqual(by_week[2], {"reps": 15, "weight": 85, "done": True})
        self.assertEqual(sorted(by_week), [1, 2])


if __name__ == "__main__":
    unittest.main()
